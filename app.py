import streamlit as st
from openpyxl import load_workbook
from io import BytesIO

st.set_page_config(
    page_title="Vindex | Preenchimento de Planilha",
    page_icon="📦",
    layout="wide"
)

st.title("📦 Vindex — Preenchimento de Quantidades")
st.write(
    "Envie a planilha Excel, selecione a aba e o sistema preencherá "
    "automaticamente as unidades a partir da coluna M, conforme a quantidade prevista na coluna J."
)

uploaded_file = st.file_uploader(
    "📂 Envie sua planilha Excel",
    type=["xlsx"],
    help="Aceita arquivos .xlsx"
)

if uploaded_file:
    try:
        file_bytes = uploaded_file.getvalue()
        wb = load_workbook(BytesIO(file_bytes))

        st.success(f"Planilha carregada: **{uploaded_file.name}**")

        sheet_name = st.selectbox(
            "📑 Escolha a aba que será preenchida",
            wb.sheetnames
        )

        ws = wb[sheet_name]

        st.info(
            "Regra: a quantidade prevista é lida da coluna **J** e o preenchimento "
            "começa na coluna **M**, colocando `1` em uma coluna para cada unidade."
        )

        if st.button("🚀 Preencher planilha", type="primary", use_container_width=True):
            total_produtos = 0
            total_unidades = 0
            linhas_processadas = []

            # J = 10
            # M = 13
            COLUNA_QUANTIDADE = 10
            COLUNA_INICIO = 13

            for row in range(1, ws.max_row + 1):
                value = ws.cell(row=row, column=COLUNA_QUANTIDADE).value

                if value is None:
                    continue

                try:
                    quantidade = int(float(value))
                except (ValueError, TypeError):
                    continue

                if quantidade <= 0:
                    continue

                # Preenche somente o intervalo necessário.
                # Não altera células fora do intervalo da quantidade.
                for i in range(quantidade):
                    ws.cell(
                        row=row,
                        column=COLUNA_INICIO + i
                    ).value = 1

                total_produtos += 1
                total_unidades += quantidade

                linhas_processadas.append((row, quantidade))

            output = BytesIO()
            wb.save(output)
            output.seek(0)

            st.success("✅ Planilha preenchida com sucesso!")

            col1, col2 = st.columns(2)
            col1.metric("Produtos processados", total_produtos)
            col2.metric("Total de unidades", total_unidades)

            if linhas_processadas:
                st.write("### 🔎 Conferência")
                for row, quantidade in linhas_processadas:
                    st.write(
                        f"• Linha **{row}** → **{quantidade}** unidades preenchidas"
                    )

            nome_original = uploaded_file.name.rsplit(".", 1)[0]
            nome_saida = f"{nome_original}_preenchida.xlsx"

            st.download_button(
                "⬇️ Baixar planilha preenchida",
                data=output.getvalue(),
                file_name=nome_saida,
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True
            )

    except Exception as e:
        st.error(f"❌ Não foi possível processar a planilha: {e}")
else:
    st.markdown(
        """
        ### Como usar

        1. Clique em **Browse files**.
        2. Selecione seu arquivo `.xlsx`.
        3. Escolha a aba.
        4. Clique em **Preencher planilha**.
        5. Baixe o novo Excel.

        **Padrão utilizado**
        - Coluna **J** → quantidade prevista
        - Coluna **M** → primeira unidade
        - Cada unidade ocupa uma coluna
        - Exemplo: quantidade 100 → M até DH recebem `1`
        """
    )
